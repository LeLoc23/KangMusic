from __future__ import annotations

import re
import textwrap
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "KangMusic_DB_Entity_Controller_View_Study_Guide.docx"


def sanitize(text: object) -> str:
    value = "" if text is None else str(text)
    value = value.replace("\t", "    ")
    chars: list[str] = []
    for ch in value:
        if ch in "\n\r":
            chars.append(ch)
            continue
        cat = unicodedata.category(ch)
        if cat in {"Cc", "Cs"}:
            continue
        chars.append(ch)
    return "".join(chars)


def set_run_font(target, name: str = "Calibri", size: float | None = None, color: str | None = None):
    font = target.font if hasattr(target, "font") else target
    font.name = name
    element = getattr(target, "_element", getattr(font, "_element", None))
    if element is not None:
        if hasattr(element, "get_or_add_rPr"):
            r_pr = element.get_or_add_rPr()
        else:
            r_pr = element
        r_fonts = getattr(r_pr, "rFonts", None)
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, r_fonts)
        r_fonts.set(qn("w:eastAsia"), name)
    if size is not None:
        font.size = Pt(size)
    if color is not None:
        font.color.rgb = RGBColor.from_string(color)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    set_run_font(normal.font, "Calibri", 11, "000000")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[name]
        set_run_font(style.font, "Calibri", size, color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    if "Source Ref" not in styles:
        style = styles.add_style("Source Ref", WD_STYLE_TYPE.PARAGRAPH)
        set_run_font(style.font, "Calibri", 8.5, "555555")
        style.font.italic = True
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(3)

    if "Code Block" not in styles:
        style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        set_run_font(style.font, "Consolas", 8, "111827")
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.0

    if "Compact Bullet" not in styles:
        style = styles.add_style("Compact Bullet", WD_STYLE_TYPE.PARAGRAPH)
        set_run_font(style.font, "Calibri", 11, "000000")
        style.base_style = styles["List Bullet"]
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    footer = section.footer.paragraphs[0]
    footer.text = "KangMusic Study Guide - DB -> Entity -> Repository/HQL -> Service -> Controller -> View"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        set_run_font(run, "Calibri", 8.5, "666666")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_in: float) -> None:
    cell.width = Inches(width_in)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_in * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = "D9E2EF", size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_paragraph(doc: Document, text: str = "", style: str | None = None, bold_prefix: str | None = None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and sanitize(text).startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        set_run_font(r1)
        r2 = p.add_run(sanitize(text)[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(sanitize(text))
        set_run_font(r)
    return p


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="Compact Bullet")
        p.add_run(sanitize(item))


def set_cell_text(cell, text: str, bold: bool = False, color: str = "000000", size: float = 10.5) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(sanitize(text))
    r.bold = bold
    set_run_font(r, "Calibri", size, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        shade_cell(hdr.cells[i], "E8EEF5")
        set_cell_text(hdr.cells[i], header, bold=True, color="0B2545", size=10)
        if widths:
            set_cell_width(hdr.cells[i], widths[i])

    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            set_cell_text(row.cells[i], value, size=9.4)
            if widths:
                set_cell_width(row.cells[i], widths[i])
    doc.add_paragraph("")
    return table


def add_callout(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, "C9D8EA")
    cell = table.cell(0, 0)
    set_cell_width(cell, 6.3)
    set_cell_margins(cell, 120, 160, 120, 160)
    shade_cell(cell, "F4F6F9")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(sanitize(title))
    r.bold = True
    set_run_font(r, "Calibri", 10.5, "0B2545")
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(sanitize(body))
    set_run_font(r2, "Calibri", 10, "111827")
    doc.add_paragraph("")


def drop_comment_noise(lines: list[tuple[int, str]]) -> list[tuple[int, str]]:
    out: list[tuple[int, str]] = []
    in_block = False
    for no, line in lines:
        stripped = line.strip()
        if stripped.startswith("/**") or stripped.startswith("/*"):
            in_block = True
            continue
        if in_block:
            if "*/" in stripped:
                in_block = False
            continue
        if stripped.startswith("*") or stripped.startswith("//"):
            continue
        if "//" in line:
            line = line.split("//", 1)[0].rstrip()
        out.append((no, line.rstrip()))
    return out


def code_snippet(path: str, start: int, end: int, skip_comments: bool = True) -> str:
    file_path = ROOT / path
    raw_lines = file_path.read_text(encoding="utf-8", errors="replace").splitlines()
    selected = [(i, raw_lines[i - 1]) for i in range(start, min(end, len(raw_lines)) + 1)]
    if skip_comments:
        selected = drop_comment_noise(selected)
    rendered: list[str] = []
    for no, line in selected:
        line = line.replace("ðŸŽµ", "[music icon]").replace("ðŸ“", "[folder icon]")
        rendered.append(f"{no:>4}: {line}")
    return "\n".join(rendered)


def manual_code(text: str) -> str:
    return textwrap.dedent(text).strip("\n")


def add_code_block(doc: Document, title: str, source: str, code: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(sanitize(title))
    r.bold = True
    set_run_font(r, "Calibri", 10.5, "1F4D78")

    s = doc.add_paragraph(style="Source Ref")
    s.add_run(sanitize(source))

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, "D9E2EF", "4")
    cell = table.cell(0, 0)
    set_cell_width(cell, 6.3)
    set_cell_margins(cell, 100, 120, 100, 120)
    shade_cell(cell, "F8FAFC")

    lines: list[str] = []
    for raw in sanitize(code).splitlines():
        if raw == "":
            lines.append("")
            continue
        wrapped = textwrap.wrap(
            raw,
            width=96,
            replace_whitespace=False,
            drop_whitespace=False,
            break_long_words=True,
            break_on_hyphens=False,
            subsequent_indent="      ",
        )
        lines.extend(wrapped or [""])

    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        p.style = "Code Block"
        run = p.add_run(line)
        set_run_font(run, "Consolas", 8, "111827")
        first = False
    doc.add_paragraph("")


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("KangMusic")
    r.bold = True
    set_run_font(r, "Calibri", 26, "0B2545")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Sổ tay vấn đáp luồng dữ liệu")
    r.bold = True
    set_run_font(r, "Calibri", 18, "2E74B5")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Database -> Entity -> Repository/HQL -> Service -> Controller -> View")
    set_run_font(r, "Calibri", 12, "333333")

    add_callout(
        doc,
        "Mục tiêu học",
        "Tài liệu này đọc project KangMusic hiện tại và gom lại theo các luồng chạy thật. "
        "Khi thầy hỏi, bạn nên trả lời theo thứ tự: bảng database nào, entity nào map bảng đó, "
        "repository/HQL nào truy vấn, service xử lý nghiệp vụ, controller đưa dữ liệu vào model/JSON, "
        "và view Thymeleaf/JavaScript hiển thị ra sao.",
    )

    add_table(
        doc,
        ["Thông tin", "Nội dung"],
        [
            ["Stack", "Spring Boot 3.2.4, Java 21, Spring MVC, Thymeleaf, Spring Data JPA, SQL Server"],
            ["Database runtime", "Dev dùng SQL Server với Hibernate ddl-auto=update; prod dùng Flyway migration và ddl-auto=validate"],
            ["Nơi query", "Chủ yếu ở repository qua method-name query và @Query JPQL/HQL"],
            ["View", "Thymeleaf templates trong src/main/resources/templates và JavaScript app.js cho AJAX/player"],
            ["Ngày lập tài liệu", "2026-06-04"],
        ],
        widths=[1.6, 4.8],
    )
    doc.add_page_break()


def add_architecture(doc: Document) -> None:
    doc.add_heading("1. Bản Đồ Kiến Trúc", level=1)
    add_callout(
        doc,
        "Câu trả lời ngắn khi thầy hỏi kiến trúc",
        "KangMusic đi theo mô hình MVC. Database SQL Server được map sang Entity JPA. "
        "Repository dùng Spring Data JPA và JPQL/HQL để truy vấn entity. Service chứa nghiệp vụ. "
        "Controller nhận request, gọi service, đưa dữ liệu vào Model hoặc JSON. View Thymeleaf và app.js render cho người dùng.",
    )
    add_code_block(
        doc,
        "Luồng đọc dữ liệu chuẩn",
        "Sơ đồ học thuộc",
        manual_code(
            """
            SQL Server table
              -> JPA Entity (@Entity, @Table, @Column, @JoinColumn)
              -> Repository (JpaRepository, method-name query, @Query JPQL/HQL)
              -> Service (@Transactional, business rules)
              -> Controller (@GetMapping/@PostMapping, Model/ResponseEntity)
              -> View (Thymeleaf th:each/th:text/th:href hoặc JavaScript fetch)
            """
        ),
    )
    add_table(
        doc,
        ["Layer", "Thư mục/file", "Vai trò khi vấn đáp"],
        [
            ["Database", "src/main/resources/db/migration", "Chứa bảng, khóa ngoại, index. Dùng để giải thích dữ liệu nằm ở đâu."],
            ["Entity", "src/main/java/com/musicapp/models", "Map bảng SQL Server sang object Java bằng JPA annotation."],
            ["Repository", "src/main/java/com/musicapp/repositories", "Nơi có HQL/JPQL @Query và method-name query."],
            ["Service", "src/main/java/com/musicapp/services", "Nơi xử lý nghiệp vụ: upload, duyệt, playlist, library, recommendation."],
            ["Controller", "src/main/java/com/musicapp/controllers", "Nơi nhận request và đưa dữ liệu sang view/API."],
            ["View", "src/main/resources/templates, static/js/app.js", "Thymeleaf render server-side; app.js gọi API cho player, library, comment."],
        ],
        widths=[1.1, 2.3, 3.0],
    )
    add_bullets(
        doc,
        [
            "Điểm cần nhớ: HQL/JPQL không viết theo tên bảng SQL như media_items mà viết theo tên entity như MediaItem.",
            "Các field trong HQL là field Java như approvalStatus, emotionLabel, uploadedAt, không phải tên cột approval_status, emotion_label, uploaded_at.",
            "Controller không nên xử lý nghiệp vụ nặng; project này phần lớn đẩy xuống Service.",
            "View không query database trực tiếp; view chỉ đọc biến model hoặc gọi API.",
        ],
    )


def add_database_entities(doc: Document) -> None:
    doc.add_heading("2. Database Và Entity", level=1)
    add_table(
        doc,
        ["Bảng SQL Server", "Entity JPA", "Ý nghĩa", "Quan hệ chính"],
        [
            ["users", "User", "Tài khoản, mật khẩu BCrypt, email, role, reset token, verify email", "1-1 creator_profiles; 1-N playlists; 1-N user_library qua user_id"],
            ["creator_profiles", "CreatorProfile", "Hồ sơ Creator/Singer chờ duyệt hoặc đã duyệt", "1-1 users; N-N media_items qua media_creators"],
            ["media_items", "MediaItem", "Bài hát/video, file, poster, lyrics, genre, playCount, trạng thái duyệt", "N-N creator_profiles; 1-N playlist_items, user_library, play_history, comments"],
            ["media_creators", "Không có entity riêng", "Bảng nối many-to-many giữa media và creator", "media_item_id + creator_profile_id"],
            ["playlists", "Playlist", "Playlist hoặc folder playlist của user", "Self parent_id; 1-N playlist_items"],
            ["playlist_items", "PlaylistItem", "Một bài trong playlist, có position", "N-1 playlists, N-1 media_items"],
            ["user_library", "UserLibrary", "Bài user đã thích/lưu", "user_id + media_item_id unique"],
            ["play_history", "PlayHistory", "Lịch sử nghe dùng cho recommendation", "N-1 media_items, user_id có thể null"],
            ["media_comments", "MediaComment", "Bình luận theo bài hát/video", "N-1 users, N-1 media_items"],
            ["song_lyrics", "SongLyrics", "Lyrics dạng text/json cho admin chỉnh và AI generate", "song_id unique trỏ media_items.id"],
        ],
        widths=[1.35, 1.35, 2.1, 1.6],
    )
    add_code_block(
        doc,
        "Entity User map bảng users",
        "src/main/java/com/musicapp/models/User.java:11-54",
        code_snippet("src/main/java/com/musicapp/models/User.java", 11, 54),
    )
    add_paragraph(
        doc,
        "Giải thích: @Entity báo đây là entity JPA; @Table(name = \"users\") map class User vào bảng users. "
        "@Id và @GeneratedValue(strategy = IDENTITY) dùng identity tự tăng của SQL Server. "
        "Các @Column unique/nullable/length giữ ràng buộc ở tầng object và giúp Hibernate tạo/validate schema.",
    )
    add_code_block(
        doc,
        "Entity MediaItem và quan hệ many-to-many creator",
        "src/main/java/com/musicapp/models/MediaItem.java:22-102",
        code_snippet("src/main/java/com/musicapp/models/MediaItem.java", 22, 102),
    )
    add_paragraph(
        doc,
        "Giải thích: MediaItem là entity quan trọng nhất. Field approvalStatus và lyricsStatus lưu enum dạng chuỗi nhờ @Enumerated(EnumType.STRING). "
        "deleted là soft delete nên nhiều query luôn thêm điều kiện m.deleted = false. "
        "Quan hệ @ManyToMany dùng @JoinTable(name = \"media_creators\") để Hibernate join sang CreatorProfile.",
    )
    add_code_block(
        doc,
        "Entity CreatorProfile map bảng creator_profiles",
        "src/main/java/com/musicapp/models/CreatorProfile.java:18-43",
        code_snippet("src/main/java/com/musicapp/models/CreatorProfile.java", 18, 43),
    )
    add_paragraph(
        doc,
        "Giải thích: CreatorProfile one-to-one với User qua user_id. Khi admin duyệt creator, CreatorService đổi status sang APPROVED và đổi role user thành ROLE_CREATOR.",
    )
    add_code_block(
        doc,
        "Playlist và PlaylistItem",
        "src/main/java/com/musicapp/models/Playlist.java:9-57; PlaylistItem.java:6-28",
        code_snippet("src/main/java/com/musicapp/models/Playlist.java", 9, 57)
        + "\n\n"
        + code_snippet("src/main/java/com/musicapp/models/PlaylistItem.java", 6, 28),
    )
    add_paragraph(
        doc,
        "Giải thích: Playlist có parent để làm folder, subPlaylists để hiển thị thư mục con, items để lấy danh sách bài. "
        "PlaylistItem là bảng trung gian có thêm position nên cần entity riêng thay vì many-to-many đơn giản.",
    )


def add_hql_section(doc: Document) -> None:
    doc.add_heading("3. Repository Và HQL/JPQL", level=1)
    add_callout(
        doc,
        "Nói đúng thuật ngữ",
        "Trong Spring Data JPA, @Query ở repository dùng JPQL/HQL kiểu entity. Hibernate là provider thực thi JPQL/HQL thành SQL Server SQL. "
        "Ngoài @Query, các hàm như findByUsername, findByIdAndDeletedFalse là derived query: Spring Data tự sinh query từ tên method.",
    )
    add_table(
        doc,
        ["Kiểu query", "Ví dụ trong project", "Cách giải thích"],
        [
            ["Derived query", "findByUsername, findByIdAndDeletedFalse, existsByUserId", "Spring Data tách tên method thành WHERE theo field entity."],
            ["JPQL/HQL SELECT", "SELECT m FROM MediaItem m WHERE m.deleted = false", "Trả về entity object, không trả ResultSet thô."],
            ["JOIN quan hệ", "LEFT JOIN m.creators c", "Join qua field object creators, không join tay bằng media_creators."],
            ["JOIN FETCH", "LEFT JOIN FETCH p.subPlaylists", "Vừa join vừa load collection trong transaction để tránh LazyInitializationException."],
            ["UPDATE/DELETE HQL", "UPDATE MediaItem m SET m.playCount = m.playCount + 1", "Cần @Modifying vì không phải SELECT."],
            ["Pageable", "PageRequest.of(0, 10)", "Không nằm trong HQL text, nhưng Spring Data thêm limit/offset khi chạy query."],
        ],
        widths=[1.3, 2.2, 2.9],
    )
    add_code_block(
        doc,
        "HQL tìm kiếm media active",
        "src/main/java/com/musicapp/repositories/MediaItemRepository.java:22-35",
        code_snippet("src/main/java/com/musicapp/repositories/MediaItemRepository.java", 22, 35),
    )
    add_paragraph(
        doc,
        "Ý nghĩa: SELECT DISTINCT m lấy MediaItem không trùng vì một bài có thể có nhiều creator. "
        "LEFT JOIN m.creators c cho phép tìm theo stageName của creator nhưng vẫn giữ bài chưa gắn creator. "
        "Điều kiện deleted=false và approvalStatus=APPROVED đảm bảo chỉ hiện bài công khai.",
    )
    add_code_block(
        doc,
        "HQL tìm bài tương tự và bài của creator",
        "src/main/java/com/musicapp/repositories/MediaItemRepository.java:70-82,105-113",
        code_snippet("src/main/java/com/musicapp/repositories/MediaItemRepository.java", 70, 82)
        + "\n\n"
        + code_snippet("src/main/java/com/musicapp/repositories/MediaItemRepository.java", 105, 113),
    )
    add_paragraph(
        doc,
        "Ý nghĩa: findSimilar lọc bài cùng genre hoặc emotionLabel và loại chính bài đang nghe bằng m.id != :excludeId. "
        "findByCreatorIdActive dùng JOIN m.creators c để lấy toàn bộ media đã duyệt thuộc một creator.",
    )
    add_code_block(
        doc,
        "HQL update lượt nghe",
        "src/main/java/com/musicapp/repositories/MediaItemRepository.java:115-118",
        code_snippet("src/main/java/com/musicapp/repositories/MediaItemRepository.java", 115, 118),
    )
    add_paragraph(
        doc,
        "Ý nghĩa: đây là bulk update HQL. Project tăng playCount trực tiếp trong database để tránh kiểu đọc playCount ra Java rồi save lại, vốn dễ bị race condition khi nhiều người nghe cùng lúc.",
    )
    add_code_block(
        doc,
        "HQL playlist: JOIN FETCH và COUNT",
        "src/main/java/com/musicapp/repositories/PlaylistRepository.java:20-32",
        code_snippet("src/main/java/com/musicapp/repositories/PlaylistRepository.java", 20, 32),
    )
    add_paragraph(
        doc,
        "Ý nghĩa: JOIN FETCH p.subPlaylists load folder con trước khi Thymeleaf render. COUNT(pi) đếm số bài mà không cần load toàn bộ PlaylistItem.",
    )
    add_code_block(
        doc,
        "HQL recommendation collaborative filtering",
        "src/main/java/com/musicapp/repositories/PlayHistoryRepository.java:59-73",
        code_snippet("src/main/java/com/musicapp/repositories/PlayHistoryRepository.java", 59, 73),
    )
    add_paragraph(
        doc,
        "Ý nghĩa: ph1 là lịch sử nghe của user có seed; ph2 là các bài khác mà những user tương tự cũng nghe. "
        "COUNT(DISTINCT ph2.userId) là điểm CF: càng nhiều user chung thì bài càng đáng gợi ý.",
    )


def add_flow_home(doc: Document) -> None:
    doc.add_heading("4. Luồng Trang Chủ, Tìm Kiếm, Lọc Genre", level=1)
    add_table(
        doc,
        ["Bước", "Thành phần", "Điều xảy ra"],
        [
            ["1", "Database", "Dữ liệu nằm ở media_items, media_creators, creator_profiles."],
            ["2", "Entity", "MediaItem map media_items; creators map many-to-many qua media_creators."],
            ["3", "Repository/HQL", "MediaItemRepository.searchActive tìm title/artist/album/creator/genre."],
            ["4", "Service", "MediaService.findPaginated gọi repository và trả Page<MediaItem>."],
            ["5", "Controller", "MusicController.homePage add mediaList, query, selectedGenre, pagination, newReleases, todayPicks."],
            ["6", "View", "index.html dùng th:each để render mediaList, todayPicks, newReleases."],
        ],
        widths=[0.55, 1.55, 4.3],
    )
    add_code_block(
        doc,
        "Service đọc trang media",
        "src/main/java/com/musicapp/services/MediaService.java:48-59",
        code_snippet("src/main/java/com/musicapp/services/MediaService.java", 48, 59),
    )
    add_code_block(
        doc,
        "Controller trang chủ đưa dữ liệu sang model",
        "src/main/java/com/musicapp/controllers/MusicController.java:77-115",
        code_snippet("src/main/java/com/musicapp/controllers/MusicController.java", 77, 115),
    )
    add_code_block(
        doc,
        "View index.html lặp mediaList",
        "src/main/resources/templates/index.html:93-149",
        code_snippet("src/main/resources/templates/index.html", 93, 149, skip_comments=False),
    )
    add_callout(
        doc,
        "Cách trả lời vấn đáp",
        "Trang chủ không query database trong HTML. Request GET / vào MusicController.homePage. Controller gọi MediaService.findPaginated. "
        "Service gọi MediaItemRepository.searchActive, HQL trả Page<MediaItem>. Controller lấy mediaPage.getContent() đưa vào model tên mediaList. "
        "index.html dùng th:each=\"s, stat : ${mediaList}\" để render từng bài.",
    )


def add_flow_track(doc: Document) -> None:
    doc.add_heading("5. Luồng Chi Tiết Bài Hát, Player, Lượt Nghe", level=1)
    add_table(
        doc,
        ["Luồng", "Database/Entity", "Controller/View"],
        [
            ["Mở /track/{id}", "media_items -> MediaItem; play_history dùng cho recommendation", "MusicController.trackDetail -> track.html"],
            ["Phát file", "fileName/posterFilename là key file đã lưu", "app.js dùng /stream/{filename}; StreamController trả ResourceRegion"],
            ["Ghi lượt nghe", "media_items.play_count + play_history", "app.js POST /api/play/{id}; RecommendationService.recordPlay"],
            ["Gợi ý", "play_history + media_items", "RecommendationService.getRecommendations rồi track.html render recommendations"],
        ],
        widths=[1.35, 2.2, 2.85],
    )
    add_code_block(
        doc,
        "Controller trang chi tiết track",
        "src/main/java/com/musicapp/controllers/MusicController.java:277-300",
        code_snippet("src/main/java/com/musicapp/controllers/MusicController.java", 277, 300),
    )
    add_code_block(
        doc,
        "View track.html đọc biến track",
        "src/main/resources/templates/track.html:53-74,158-160",
        code_snippet("src/main/resources/templates/track.html", 53, 74, skip_comments=False)
        + "\n\n"
        + code_snippet("src/main/resources/templates/track.html", 158, 160, skip_comments=False),
    )
    add_code_block(
        doc,
        "Ghi lượt nghe: controller -> service -> HQL update",
        "MusicController.java:209-214; RecommendationService.java:113-124; MediaItemRepository.java:115-118",
        code_snippet("src/main/java/com/musicapp/controllers/MusicController.java", 209, 214)
        + "\n\n"
        + code_snippet("src/main/java/com/musicapp/services/RecommendationService.java", 113, 124)
        + "\n\n"
        + code_snippet("src/main/java/com/musicapp/repositories/MediaItemRepository.java", 115, 118),
    )
    add_code_block(
        doc,
        "Streaming file có hỗ trợ byte range",
        "src/main/java/com/musicapp/controllers/StreamController.java:48-104",
        code_snippet("src/main/java/com/musicapp/controllers/StreamController.java", 48, 104),
    )
    add_callout(
        doc,
        "Cách trả lời vấn đáp",
        "Khi mở chi tiết bài, controller dùng mediaService.findById để chỉ lấy bài chưa deleted và đã APPROVED. "
        "View nhận biến track, similar, recommendations. Khi người dùng bấm play, JavaScript gọi POST /api/play/{id}; service tăng playCount bằng HQL UPDATE và nếu user đăng nhập thì thêm một bản ghi PlayHistory để phục vụ recommendation.",
    )


def add_flow_upload(doc: Document) -> None:
    doc.add_heading("6. Luồng Upload, Edit, Duyệt Media", level=1)
    add_table(
        doc,
        ["Bước", "File/code", "Giải thích"],
        [
            ["1", "add.html form", "Form multipart gửi title, artist, mediaFile, posterFile, genre, album, type, creatorIds."],
            ["2", "MusicController.saveMedia", "Nhận request, xác định admin hay creator, gom creatorIds, gọi MediaService.saveMedia."],
            ["3", "MediaService.saveMedia", "Kiểm tra extension, lưu file, tạo MediaItem, set status APPROVED nếu admin hoặc PENDING nếu creator."],
            ["4", "Repository.save", "Hibernate insert media_items và bảng nối media_creators nếu có creator."],
            ["5", "Admin duyệt", "AdminController gọi mediaService.approveMedia/rejectMedia, cập nhật approvalStatus."],
            ["6", "View", "creator.html/admin.html hiển thị trạng thái pending/approved/rejected."],
        ],
        widths=[0.55, 1.75, 4.1],
    )
    add_code_block(
        doc,
        "Form upload ở view",
        "src/main/resources/templates/add.html:36-58",
        code_snippet("src/main/resources/templates/add.html", 36, 58, skip_comments=False),
    )
    add_code_block(
        doc,
        "Controller nhận POST /add",
        "src/main/java/com/musicapp/controllers/MusicController.java:133-164",
        code_snippet("src/main/java/com/musicapp/controllers/MusicController.java", 133, 164),
    )
    add_code_block(
        doc,
        "Service tạo MediaItem và lưu database",
        "src/main/java/com/musicapp/services/MediaService.java:112-154",
        code_snippet("src/main/java/com/musicapp/services/MediaService.java", 112, 154),
    )
    add_code_block(
        doc,
        "Admin approve/reject media",
        "src/main/java/com/musicapp/controllers/AdminController.java:107-113; MediaService.java:209-238",
        code_snippet("src/main/java/com/musicapp/controllers/AdminController.java", 107, 113)
        + "\n\n"
        + code_snippet("src/main/java/com/musicapp/services/MediaService.java", 209, 238),
    )
    add_callout(
        doc,
        "Cách trả lời vấn đáp",
        "Upload không insert SQL thủ công. Service tạo object MediaItem rồi gọi mediaItemRepository.save(item). "
        "Nếu người upload là admin thì approvalStatus = APPROVED; nếu creator thì PENDING để admin duyệt. "
        "View chỉ hiện những bài đã APPROVED vì các HQL public đều có điều kiện approvalStatus APPROVED.",
    )


def add_flow_playlist(doc: Document) -> None:
    doc.add_heading("7. Luồng Playlist Và Folder", level=1)
    add_table(
        doc,
        ["Luồng", "Repository/HQL", "View"],
        [
            ["Danh sách playlist sidebar", "findRootPlaylistsWithChildren dùng LEFT JOIN FETCH p.subPlaylists", "GlobalModelAdvice thêm globalPlaylists cho mọi trang Thymeleaf"],
            ["Xem playlist", "playlistRepo.findById; entity lazy load items trong transaction/view", "playlist.html dùng playlist.items và pi.mediaItem"],
            ["Tạo playlist", "playlistRepo.save(new Playlist(...))", "Form/JS POST /playlists/create"],
            ["Thêm bài", "existsByPlaylistIdAndMediaItemId, findMaxPosition, save PlaylistItem", "Modal Add to playlist trong app.js"],
            ["Xóa item", "playlistItemRepo.delete", "playlist.html form /playlists/item/{itemId}/remove"],
        ],
        widths=[1.55, 2.7, 2.15],
    )
    add_code_block(
        doc,
        "Service load playlist tránh LazyInitializationException",
        "src/main/java/com/musicapp/services/PlaylistService.java:35-48; PlaylistRepository.java:20",
        code_snippet("src/main/java/com/musicapp/services/PlaylistService.java", 35, 48)
        + "\n\n"
        + code_snippet("src/main/java/com/musicapp/repositories/PlaylistRepository.java", 20, 20),
    )
    add_code_block(
        doc,
        "Thêm bài vào playlist",
        "src/main/java/com/musicapp/services/PlaylistService.java:107-125",
        code_snippet("src/main/java/com/musicapp/services/PlaylistService.java", 107, 125),
    )
    add_code_block(
        doc,
        "HQL kiểm tra và lấy position",
        "src/main/java/com/musicapp/repositories/PlaylistItemRepository.java:15-23",
        code_snippet("src/main/java/com/musicapp/repositories/PlaylistItemRepository.java", 15, 23),
    )
    add_code_block(
        doc,
        "Controller view playlist và view render items",
        "PlaylistController.java:46-64; playlist.html:101-171",
        code_snippet("src/main/java/com/musicapp/controllers/PlaylistController.java", 46, 64)
        + "\n\n"
        + code_snippet("src/main/resources/templates/playlist.html", 101, 171, skip_comments=False),
    )
    add_callout(
        doc,
        "Cách trả lời vấn đáp",
        "PlaylistItem là entity riêng vì bảng playlist_items không chỉ nối playlist với media mà còn có position và addedAt. "
        "Khi thêm bài, service kiểm tra owner qua findByIdAndUserId, chống trùng bằng existsByPlaylistIdAndMediaItemId, lấy MAX(position) rồi save PlaylistItem mới.",
    )


def add_flow_library_comments(doc: Document) -> None:
    doc.add_heading("8. Luồng Thư Viện Và Bình Luận", level=1)
    doc.add_heading("8.1 Thư viện yêu thích", level=2)
    add_table(
        doc,
        ["Bước", "Code", "Giải thích"],
        [
            ["DB", "user_library", "Mỗi dòng là một bài user đã thích; unique user_id + media_item_id chống trùng."],
            ["Entity", "UserLibrary", "Lưu userId và ManyToOne mediaItem."],
            ["Repository", "findByUserIdOrderByAddedAtDesc, findLikedMediaIds", "Lấy danh sách thư viện và set id để highlight nút tim."],
            ["Service", "LibraryService.toggleLibrary", "Nếu đã có thì delete; chưa có thì save UserLibrary mới."],
            ["Controller/View", "LibraryController + library.html + app.js", "GET /library render danh sách; POST /api/library/toggle dùng AJAX."],
        ],
        widths=[0.75, 2.2, 3.45],
    )
    add_code_block(
        doc,
        "LibraryService toggle",
        "src/main/java/com/musicapp/services/LibraryService.java:29-61",
        code_snippet("src/main/java/com/musicapp/services/LibraryService.java", 29, 61),
    )
    add_code_block(
        doc,
        "LibraryController và view",
        "LibraryController.java:30-56; library.html:67-111",
        code_snippet("src/main/java/com/musicapp/controllers/LibraryController.java", 30, 56)
        + "\n\n"
        + code_snippet("src/main/resources/templates/library.html", 67, 111, skip_comments=False),
    )
    doc.add_heading("8.2 Bình luận theo media", level=2)
    add_code_block(
        doc,
        "CommentController đọc và ghi comment",
        "src/main/java/com/musicapp/controllers/CommentController.java:31-84",
        code_snippet("src/main/java/com/musicapp/controllers/CommentController.java", 31, 84),
    )
    add_paragraph(
        doc,
        "Giải thích: comment API trả JSON chứ không render Thymeleaf trực tiếp. app.js gọi /api/comments/{mediaId}, nhận danh sách map gồm id, content, timestampSeconds, createdAt và user displayName để render ở khu bình luận của track.",
    )


def add_flow_user_creator_admin_lyrics(doc: Document) -> None:
    doc.add_heading("9. Luồng User, Creator, Admin, Lyrics", level=1)
    doc.add_heading("9.1 Đăng ký, đăng nhập, profile", level=2)
    add_code_block(
        doc,
        "AuthService đăng ký user",
        "src/main/java/com/musicapp/services/AuthService.java:38-66",
        code_snippet("src/main/java/com/musicapp/services/AuthService.java", 38, 66),
    )
    add_paragraph(
        doc,
        "Giải thích: đăng ký kiểm tra trống, độ dài, username/email trùng, độ mạnh mật khẩu, xác nhận mật khẩu. Sau đó passwordEncoder.encode(password) rồi save User. "
        "Email verification code được lưu vào users.email_verification_code và expiry.",
    )
    doc.add_heading("9.2 Xin làm Creator/Singer và admin duyệt", level=2)
    add_code_block(
        doc,
        "Profile request creator và CreatorService approve",
        "ProfileController.java:54-61; CreatorService.java:24-52,54-64",
        code_snippet("src/main/java/com/musicapp/controllers/ProfileController.java", 54, 61)
        + "\n\n"
        + code_snippet("src/main/java/com/musicapp/services/CreatorService.java", 24, 52)
        + "\n\n"
        + code_snippet("src/main/java/com/musicapp/services/CreatorService.java", 54, 64),
    )
    add_paragraph(
        doc,
        "Giải thích: user gửi request ở /profile/creator/request. Service tạo hoặc cập nhật CreatorProfile status PENDING. Admin duyệt thì status APPROVED và role của User đổi thành ROLE_CREATOR.",
    )
    doc.add_heading("9.3 Admin dashboard", level=2)
    add_code_block(
        doc,
        "Admin dashboard lấy dữ liệu",
        "src/main/java/com/musicapp/controllers/AdminController.java:35-41; admin.html snippets",
        code_snippet("src/main/java/com/musicapp/controllers/AdminController.java", 35, 41)
        + "\n\n"
        + code_snippet("src/main/resources/templates/admin.html", 34, 45, skip_comments=False)
        + "\n\n"
        + code_snippet("src/main/resources/templates/admin.html", 74, 86, skip_comments=False),
    )
    add_paragraph(
        doc,
        "Giải thích: admin.html có nhiều bảng: creatorRequests, pendingMedia, users, mediaItems. Dữ liệu đều được AdminController addAttribute từ các service.",
    )
    doc.add_heading("9.4 Lyrics", level=2)
    add_code_block(
        doc,
        "Lyrics save và async generate",
        "AdminLyricsController.java:86-121; LyricsService.java:40-54",
        code_snippet("src/main/java/com/musicapp/controllers/AdminLyricsController.java", 86, 121)
        + "\n\n"
        + code_snippet("src/main/java/com/musicapp/services/LyricsService.java", 40, 54),
    )
    add_paragraph(
        doc,
        "Giải thích: bảng song_lyrics lưu lyricsText và lyricsJson. Khi admin save lyrics, vừa save SongLyrics vừa đồng bộ MediaItem.lyrics và lyricsStatus. Khi generate, LyricsService chạy @Async, đổi status PROCESSING rồi xử lý audio/Whisper và lưu kết quả.",
    )


def add_route_map(doc: Document) -> None:
    doc.add_heading("10. Bảng Route -> Controller -> DB/View", level=1)
    add_table(
        doc,
        ["Route", "Controller method", "Service/Repository chính", "View/Response"],
        [
            ["/", "MusicController.homePage", "MediaService.findPaginated -> MediaItemRepository.searchActive", "index.html"],
            ["/track/{id}", "MusicController.trackDetail", "MediaService.findById, findSimilar, RecommendationService", "track.html"],
            ["/api/play/{id}", "MusicController.recordPlay", "RecommendationService.recordPlay -> incrementPlayCount, PlayHistory save", "200 OK"],
            ["/stream/{filename}", "StreamController.stream", "File system từ app.upload.dir", "ResourceRegion audio/video"],
            ["/add GET/POST", "MusicController.showAddForm/saveMedia", "CreatorService, MediaService.saveMedia", "add.html hoặc redirect"],
            ["/media/{id}/edit", "MusicController.editMediaForm/updateMedia", "MediaService.findByIdForManagement/updateMedia", "edit-media.html"],
            ["/library", "LibraryController.libraryPage", "LibraryService.getUserLibrary", "library.html"],
            ["/api/library/toggle", "LibraryController.toggleLike", "LibraryService.toggleLibrary", "JSON liked"],
            ["/playlists/{id}", "PlaylistController.viewPlaylist", "PlaylistService.getPlaylistById", "playlist.html"],
            ["/playlists/{id}/add", "PlaylistController.addToPlaylist", "PlaylistService.addToPlaylist", "JSON success"],
            ["/api/comments/{mediaId}", "CommentController", "MediaCommentRepository, MediaItemRepository", "JSON comments"],
            ["/profile", "ProfileController.profilePage", "UserService, CreatorService", "profile.html"],
            ["/profile/creator/request", "ProfileController.requestCreator", "CreatorService.requestCreator", "redirect profile"],
            ["/creator", "CreatorController.creatorDashboard", "CreatorService, MediaService.findByUploader", "creator.html"],
            ["/admin", "AdminController.adminDashboard", "UserService, MediaService, CreatorService", "admin.html"],
            ["/admin/songs/{id}/lyrics/edit", "AdminLyricsController.editLyricsForm", "MediaItemRepository, SongLyricsRepository", "admin-lyrics-edit.html"],
            ["/api/ai/recommend", "AiRecommendationController", "MediaService.findAllActiveList, OpenAiService", "JSON tracks"],
        ],
        widths=[1.45, 1.65, 2.25, 1.05],
    )
    doc.add_heading("Security route cần nhớ", level=2)
    add_code_block(
        doc,
        "SecurityConfig phân quyền",
        "src/main/java/com/musicapp/config/SecurityConfig.java:47-67",
        code_snippet("src/main/java/com/musicapp/config/SecurityConfig.java", 47, 67),
    )
    add_paragraph(
        doc,
        "Giải thích: /admin/** cần ROLE_ADMIN; /creator, /add, /media/*/edit cần authenticated; /playlists/** và /library cần authenticated. "
        "Một số route đọc như /track/**, /api/search, /api/genre/**, /api/similar/** được public.",
    )


def add_hql_oral_cheatsheet(doc: Document) -> None:
    doc.add_heading("11. HQL/JPQL Vấn Đáp", level=1)
    add_table(
        doc,
        ["Câu hỏi", "Câu trả lời nên nói"],
        [
            ["HQL nằm ở đâu?", "Trong các repository có @Query, ví dụ MediaItemRepository, PlaylistRepository, PlayHistoryRepository, UserLibraryRepository, PlaylistItemRepository."],
            ["HQL khác SQL thế nào?", "HQL query theo entity và field Java: SELECT m FROM MediaItem m, không SELECT * FROM media_items."],
            ["Vì sao dùng SELECT DISTINCT m?", "Khi join m.creators, một media có nhiều creator có thể bị lặp, DISTINCT loại trùng."],
            ["LEFT JOIN m.creators c để làm gì?", "Để tìm theo creator.stageName nhưng vẫn giữ bài không có creator; nếu JOIN thường thì bài không có creator có thể mất."],
            ["JOIN FETCH khác JOIN thường?", "JOIN FETCH vừa join vừa load collection vào entity, tránh lỗi lazy khi view render ngoài transaction."],
            ["@Modifying dùng khi nào?", "Dùng cho @Query UPDATE/DELETE, ví dụ incrementPlayCount và deleteOlderThan."],
            ["Pageable hoạt động ra sao?", "Repository nhận Pageable, Spring Data thêm limit/offset và sort vào SQL sinh ra từ JPQL."],
            ["Derived query là gì?", "Method như findByUsername hoặc findByIdAndDeletedFalse tự được Spring Data parse thành query."],
            ["Soft delete nằm ở đâu?", "MediaItem có field deleted; query public luôn lọc m.deleted = false."],
            ["Duyệt media nằm ở đâu?", "MediaItem.approvalStatus; upload bởi creator là PENDING, admin approve đổi thành APPROVED."],
            ["Tại sao không expose entity trực tiếp ở API search?", "Dùng MediaItemDto để trả src public /stream/{uuid} và không lộ raw fileName/deleted/internal fields."],
            ["Gợi ý nhạc dùng query nào?", "PlayHistoryRepository.findRecommendationsFromSeeds join PlayHistory ph1 với ph2 theo userId, group by media và count user chung."],
            ["Tăng lượt nghe an toàn hơn ở đâu?", "MediaItemRepository.incrementPlayCount dùng HQL UPDATE trực tiếp nên tránh race condition đọc-sửa-ghi."],
            ["Playlist chống thêm trùng thế nào?", "PlaylistItemRepository.existsByPlaylistIdAndMediaItemId kiểm tra trước khi save PlaylistItem."],
            ["Bảng media_creators có entity riêng không?", "Không. Nó được dùng bởi @ManyToMany trong MediaItem thông qua @JoinTable."],
            ["PlaylistItem vì sao có entity riêng?", "Vì bảng playlist_items có dữ liệu phụ như position, addedAt; không chỉ là bảng nối đơn giản."],
            ["@Transactional để làm gì?", "Đặt biên giao dịch cho service; readOnly=true tối ưu đọc; write transaction đảm bảo save/update cùng phiên làm việc."],
            ["FetchType.LAZY/EAGER ví dụ?", "Playlist.items LAZY để tránh load nặng; PlaylistItem.mediaItem EAGER để khi render item có sẵn media."],
            ["Controller có query DB trực tiếp không?", "Một vài controller như AdminLyricsController dùng repository trực tiếp, nhưng luồng chính thường đi qua service."],
            ["View lấy dữ liệu như thế nào?", "Controller addAttribute vào Model, Thymeleaf dùng th:each/th:text/th:href; AJAX thì controller trả JSON cho app.js."],
        ],
        widths=[2.0, 4.4],
    )
    add_code_block(
        doc,
        "Mẫu trả lời HQL searchActive theo từng dòng",
        "Ghi nhớ khi thầy chỉ vào query MediaItemRepository.searchActive",
        manual_code(
            """
            SELECT DISTINCT m FROM MediaItem m
              -> lấy entity MediaItem, đặt bí danh m, DISTINCT để không lặp khi join creator.

            LEFT JOIN m.creators c
              -> join qua field creators của entity, không join bảng media_creators bằng SQL tay.

            WHERE m.deleted = false
              -> bỏ media soft-deleted.

            AND m.approvalStatus = MediaApprovalStatus.APPROVED
              -> chỉ hiện media đã được duyệt công khai.

            LOWER(m.title) LIKE LOWER(CONCAT('%', :query, '%'))
              -> tìm không phân biệt hoa thường theo title; :query là tham số từ controller.

            AND (:genre IS NULL OR :genre = '' OR m.genre = :genre)
              -> nếu không chọn genre thì bỏ qua filter; nếu có thì lọc đúng genre.

            ORDER BY m.id DESC
              -> bài mới hơn có id lớn hơn sẽ đứng trước.
            """
        ),
    )
    add_code_block(
        doc,
        "Mẫu trả lời HQL collaborative filtering",
        "Ghi nhớ khi thầy hỏi query PlayHistoryRepository.findRecommendationsFromSeeds",
        manual_code(
            """
            FROM PlayHistory ph1 JOIN PlayHistory ph2 ON ph1.userId = ph2.userId
              -> tự join bảng lịch sử nghe theo cùng userId.

            WHERE ph1.mediaItem.id IN :seedIds
              -> ph1 đại diện các bài user hiện tại đã nghe/thích, gọi là seed.

            AND ph2.mediaItem.id NOT IN :seedIds
              -> ph2 là bài đề xuất, không trùng bài seed.

            AND ph2.userId != :userId
              -> không dùng chính lịch sử của user hiện tại để tự đề xuất lại.

            GROUP BY ph2.mediaItem
              -> gom theo từng bài ứng viên.

            COUNT(DISTINCT ph2.userId) AS score
              -> điểm càng cao nếu nhiều user khác cùng nghe bài đó.

            ORDER BY score DESC
              -> bài có nhiều user chung đứng trước.
            """
        ),
    )


def add_quick_study_plan(doc: Document) -> None:
    doc.add_heading("12. Cách Học Thuộc Trước Khi Vấn Đáp", level=1)
    add_table(
        doc,
        ["Thời gian", "Việc cần học", "Kết quả cần nói được"],
        [
            ["10 phút", "Học sơ đồ tổng: DB -> Entity -> Repository/HQL -> Service -> Controller -> View", "Nói được kiến trúc MVC và vai trò từng layer."],
            ["20 phút", "Học 5 entity chính: User, MediaItem, CreatorProfile, Playlist, PlaylistItem", "Nói được bảng nào map entity nào và quan hệ chính."],
            ["30 phút", "Học 4 query HQL: searchActive, findSimilar, JOIN FETCH playlist, collaborative filtering", "Giải thích được SELECT/FROM/JOIN/WHERE/ORDER BY theo entity."],
            ["20 phút", "Học 4 luồng: trang chủ, track/play, upload/approve, playlist/library", "Nói theo thứ tự request vào controller nào, gọi service nào, view nào render."],
            ["10 phút", "Ôn security và DTO", "Nói được route nào cần role, vì sao API dùng DTO."],
        ],
        widths=[1.0, 2.6, 2.8],
    )
    add_callout(
        doc,
        "Mẫu mở đầu khi bị hỏi bất kỳ chức năng nào",
        "Em sẽ lần theo luồng từ database lên view: bảng SQL Server lưu dữ liệu, entity JPA map bảng, repository dùng Spring Data JPA/HQL truy vấn entity, service xử lý nghiệp vụ trong transaction, controller nhận request và đưa dữ liệu vào Model hoặc JSON, cuối cùng Thymeleaf/app.js render ra giao diện.",
    )


def build() -> None:
    doc = Document()
    configure_document(doc)
    add_title_page(doc)
    add_architecture(doc)
    add_database_entities(doc)
    add_hql_section(doc)
    add_flow_home(doc)
    add_flow_track(doc)
    add_flow_upload(doc)
    add_flow_playlist(doc)
    add_flow_library_comments(doc)
    add_flow_user_creator_admin_lyrics(doc)
    add_route_map(doc)
    add_hql_oral_cheatsheet(doc)
    add_quick_study_plan(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
